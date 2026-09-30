"""DOCX signature placement. No document or signature is stored on disk."""

from __future__ import annotations

import base64
import binascii
from dataclasses import dataclass
from io import BytesIO
import re
from zipfile import ZipFile

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from PIL import Image, UnidentifiedImageError

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_EXPANDED_BYTES = 80 * 1024 * 1024


class SigningError(ValueError):
    """An error safe to display directly to the person signing."""


@dataclass(frozen=True)
class SignatureLocation:
    index: int
    label: str
    context: str
    place: str


def load_document(content: bytes) -> DocumentType:
    if not content or len(content) > MAX_UPLOAD_BYTES:
        raise SigningError("Cargá un documento Word de hasta 10 MB.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(e.file_size for e in entries) > MAX_EXPANDED_BYTES:
                raise SigningError("El contenido del Word es demasiado grande para procesarlo.")
            if any(e.flag_bits & 1 for e in entries):
                raise SigningError("Quitá la contraseña del Word antes de cargarlo.")
            if "word/document.xml" not in archive.namelist():
                raise SigningError("El archivo no es un documento Word .docx válido.")
        return Document(BytesIO(content))
    except SigningError:
        raise
    except Exception as exc:
        # python-docx can raise different errors for malformed XML/relationships.
        raise SigningError("No se pudo abrir el Word. Abrilo en Word y guardá una copia .docx sin contraseña.") from exc


def paragraphs_in_order(parent, place="Documento"):
    """Walk body paragraphs and nested tables; merged cells occur only once."""
    for block in parent.iter_inner_content():
        if isinstance(block, Paragraph):
            yield block, place
        elif isinstance(block, Table):
            seen = set()
            for row_number, row in enumerate(block.rows, 1):
                for column_number, cell in enumerate(row.cells, 1):
                    if cell._tc in seen:
                        continue
                    seen.add(cell._tc)
                    yield from paragraphs_in_order(
                        cell, f"{place} · tabla, fila {row_number}, columna {column_number}"
                    )


def normalize(text: str) -> str:
    return " ".join(text.casefold().split())


def find_locations(document: DocumentType, phrase: str = "firma") -> list[SignatureLocation]:
    needle = normalize(phrase)
    if not needle:
        return []
    # Match whole words (underscores may be the signature's printed line).
    pattern = re.compile(r"(?<![^\W_])" + re.escape(needle) + r"(?![^\W_])")
    paragraphs = list(paragraphs_in_order(document))
    locations = []
    for index, (paragraph, place) in enumerate(paragraphs):
        if pattern.search(normalize(paragraph.text)):
            context = "\n".join(
                p.text.strip() for p, _ in paragraphs[max(0, index - 1):index + 2] if p.text.strip()
            )
            locations.append(SignatureLocation(index, paragraph.text.strip(), context, place))
    return locations


def signature_png(data_url: str | None) -> bytes:
    if not isinstance(data_url, str) or not data_url.startswith("data:image/png;base64,"):
        raise SigningError("Dibujá tu firma antes de preparar el documento.")
    if len(data_url) > 3_000_000:
        raise SigningError("La imagen de la firma es demasiado grande. Borrala y volvé a dibujarla.")
    try:
        raw = base64.b64decode(data_url.split(",", 1)[1], validate=True)
        with Image.open(BytesIO(raw)) as image:
            if image.format != "PNG" or image.width * image.height > 4_000_000:
                raise SigningError("La imagen de la firma no es válida.")
            rgba = image.convert("RGBA")
        bbox = rgba.getchannel("A").getbbox()
        if not bbox:
            raise SigningError("El recuadro está vacío. Dibujá tu firma.")
        cropped = rgba.crop(bbox)
        padded = Image.new("RGBA", (cropped.width + 12, cropped.height + 12))
        padded.paste(cropped, (6, 6))
        output = BytesIO()
        padded.save(output, format="PNG")
        return output.getvalue()
    except SigningError:
        raise
    except (ValueError, binascii.Error, OSError, UnidentifiedImageError, Image.DecompressionBombError) as exc:
        raise SigningError("No se pudo leer la firma. Borrala y volvé a dibujarla.") from exc


def effective_format(paragraph: Paragraph, name: str):
    value = getattr(paragraph.paragraph_format, name)
    style = paragraph.style
    while value is None and style is not None:
        value = getattr(style.paragraph_format, name)
        style = style.base_style
    return value


def available_width(document: DocumentType, paragraph: Paragraph) -> int:
    # Use the section ending at/after this body block, including table blocks.
    block = paragraph._p
    while block.getparent() is not document._element.body:
        block = block.getparent()
    section_index = 0
    for child in document._element.body:
        if child is block:
            break
        if child.tag == qn("w:p") and child.find("w:pPr/w:sectPr", child.nsmap) is not None:
            section_index += 1
    section = document.sections[min(section_index, len(document.sections) - 1)]
    width = int(section.page_width - section.left_margin - section.right_margin)
    cols = section._sectPr.find(qn("w:cols"))
    if cols is not None:
        explicit = cols.findall(qn("w:col"))
        if explicit:
            width = min(width, min(int(c.get(qn("w:w"))) * 635 for c in explicit))
        else:
            count = int(cols.get(qn("w:num"), "1"))
            gap = int(cols.get(qn("w:space"), "720")) * 635
            width = min(width, (width - (count - 1) * gap) // count)
    if isinstance(paragraph._parent, _Cell):
        cell = paragraph._parent
        if cell.width is not None:
            # Reserve the cell's actual side margins, or Word's usual 108 twips.
            margins = cell._tc.tcPr.find(qn("w:tcMar"))
            table = cell._tc.getparent().getparent()
            table_margins = table.find("w:tblPr/w:tblCellMar", table.nsmap)
            padding = 0
            for side, alternate in (("left", "start"), ("right", "end")):
                node = None
                for source in (margins, table_margins):
                    if source is not None:
                        node = source.find(qn(f"w:{side}"))
                        if node is None:
                            node = source.find(qn(f"w:{alternate}"))
                        if node is not None:
                            break
                padding += int(node.get(qn("w:w"), "108")) * 635 if node is not None else 108 * 635
            width = min(width, int(cell.width) - padding)
    for side in ("left_indent", "right_indent"):
        width -= max(0, effective_format(paragraph, side) or 0)
    if width <= 0:
        raise SigningError("No hay espacio horizontal suficiente junto a ese rótulo.")
    return width


def sign_document(original: bytes, location_index: int, png: bytes) -> bytes:
    document = load_document(original)
    paragraphs = list(paragraphs_in_order(document))
    if not 0 <= location_index < len(paragraphs):
        raise SigningError("Elegí de nuevo el lugar donde debe ir la firma.")
    label, _ = paragraphs[location_index]
    with Image.open(BytesIO(png)) as image:
        ratio = image.width / image.height
    width = min(int(Cm(5)), available_width(document, label), int(Cm(2) * ratio))
    height = int(width / ratio)
    signature = label.insert_paragraph_before()
    signature.alignment = effective_format(label, "alignment")
    fmt = signature.paragraph_format
    fmt.left_indent = effective_format(label, "left_indent")
    fmt.right_indent = effective_format(label, "right_indent")
    fmt.first_line_indent = Pt(0)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1.0
    fmt.keep_with_next = True
    fmt.keep_together = True
    # A page break on the label must precede the signature as well.
    if effective_format(label, "page_break_before"):
        fmt.page_break_before = True
        label.paragraph_format.page_break_before = False
    # Also handle a manual page break directly before the label's first text.
    for child in label._p.iter():
        if child.tag == qn("w:t") and (child.text or "").strip():
            break
        if child.tag == qn("w:br") and child.get(qn("w:type")) == "page":
            child.getparent().remove(child)
            signature.add_run().add_break(WD_BREAK.PAGE)
    picture = signature.add_run().add_picture(BytesIO(png), width=width, height=height)
    picture._inline.docPr.set("descr", "Firma manuscrita")
    output = BytesIO()
    document.save(output)
    return output.getvalue()
