import base64
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Cm
from PIL import Image, ImageDraw
import pytest

from signing import (
    MAX_UPLOAD_BYTES, SigningError, available_width, find_locations,
    load_document, paragraphs_in_order, sign_document, signature_png,
)


def doc_bytes(document):
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def drawn_signature(empty=False):
    image = Image.new("RGBA", (400, 160))
    if not empty:
        ImageDraw.Draw(image).line([(40, 100), (130, 35), (100, 120), (290, 60)], fill="#173b69", width=3)
    stream = BytesIO()
    image.save(stream, format="PNG")
    return "data:image/png;base64," + base64.b64encode(stream.getvalue()).decode()


def test_generic_search_and_split_formatting():
    document = Document()
    document.add_paragraph("Confirmación: debe firmarse el documento.")
    label = document.add_paragraph()
    label.add_run("FIR").bold = True
    label.add_run("MA\u00a0 del   responsable")
    document.add_paragraph("___Firma___")
    assert len(find_locations(document)) == 2
    assert len(find_locations(document, "firma del responsable")) == 1
    assert not find_locations(document, "docente")
    assert not find_locations(document, "  ")


def test_tables_nested_and_merged_cells_are_not_duplicated():
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Firma de la persona"
    nested = table.cell(1, 1).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Firma del responsable"
    found = find_locations(document)
    assert len(found) == 2
    signed = load_document(sign_document(doc_bytes(document), found[1].index, signature_png(drawn_signature())))
    inner_cell = signed.tables[0].cell(1, 1).tables[0].cell(0, 0)
    assert inner_cell.paragraphs[0]._p.xpath(".//w:drawing")
    assert inner_cell.paragraphs[1].text == "Firma del responsable"
    assert not signed.tables[0].cell(0, 0)._tc.xpath(".//w:drawing")


def test_insertion_keeps_label_format_and_original_and_does_not_duplicate():
    document = Document()
    document.add_paragraph("Contenido original")
    label = document.add_paragraph("Firma")
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label.runs[0].bold = True
    original = doc_bytes(document)
    png = signature_png(drawn_signature())
    outputs = [sign_document(original, 1, png) for _ in range(2)]
    for content in outputs:
        signed = load_document(content)
        assert len(signed.inline_shapes) == 1
        assert signed.paragraphs[0].text == "Contenido original"
        assert signed.paragraphs[1].alignment == WD_ALIGN_PARAGRAPH.CENTER
        assert signed.paragraphs[1].paragraph_format.keep_with_next is True
        assert signed.paragraphs[2].text == "Firma"
        assert signed.paragraphs[2].runs[0].bold is True
        shape = signed.inline_shapes[0]
        assert shape.width <= Cm(5)
        assert shape.height <= Cm(2)
        with Image.open(BytesIO(png)) as image:
            assert shape.width / shape.height == pytest.approx(image.width / image.height, abs=0.0001)
    assert len(load_document(original).inline_shapes) == 0


def test_narrow_table_respects_cell_width():
    document = Document()
    table = document.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Cm(2)
    cell = table.cell(0, 0)
    cell.width = Cm(2)
    cell.text = "Firma"
    width = available_width(document, cell.paragraphs[0])
    result = load_document(sign_document(doc_bytes(document), 0, signature_png(drawn_signature())))
    assert result.inline_shapes[0].width <= width < Cm(2)


@pytest.mark.parametrize("manual", [False, True])
def test_page_break_moves_before_signature(manual):
    document = Document()
    document.add_paragraph("Página anterior")
    label = document.add_paragraph()
    if manual:
        label.add_run().add_break(WD_BREAK.PAGE)
    else:
        label.paragraph_format.page_break_before = True
    label.add_run("Firma")
    signed = load_document(sign_document(doc_bytes(document), 1, signature_png(drawn_signature())))
    signature, label = signed.paragraphs[1:3]
    assert signature.paragraph_format.keep_with_next
    assert signature.paragraph_format.page_break_before or signature._p.xpath('.//w:br[@w:type="page"]')
    assert not label.paragraph_format.page_break_before
    assert not label._p.xpath('.//w:br[@w:type="page"]')


def test_signature_is_cropped_and_transparent():
    png = signature_png(drawn_signature())
    with Image.open(BytesIO(png)) as image:
        assert image.width < 400
        assert image.height < 160
        assert image.mode == "RGBA"
        assert image.getpixel((0, 0))[3] == 0


@pytest.mark.parametrize("value", [None, "bad", "data:image/png;base64,xxx", drawn_signature(empty=True)], ids=["missing", "wrong-type", "bad-base64", "empty-canvas"])
def test_invalid_or_empty_signatures(value):
    with pytest.raises(SigningError):
        signature_png(value)


@pytest.mark.parametrize("value", [b"", b"invalid", b"x" * (MAX_UPLOAD_BYTES + 1)], ids=["empty", "corrupt", "too-large"])
def test_invalid_documents(value):
    with pytest.raises(SigningError):
        load_document(value)


def test_non_word_zip():
    stream = BytesIO()
    with ZipFile(stream, "w") as archive:
        archive.writestr("notes.txt", "not a Word document")
    with pytest.raises(SigningError):
        load_document(stream.getvalue())


def test_headers_and_footers_are_not_searched():
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Firma"
    document.sections[0].footer.paragraphs[0].text = "Firma"
    document.add_paragraph("Texto normal")
    assert not find_locations(document)


def test_independent_documents():
    first, second = Document(), Document()
    first.add_paragraph("Firma de Ana")
    second.add_paragraph("Firma de Juan")
    signed = load_document(sign_document(doc_bytes(first), 0, signature_png(drawn_signature())))
    assert signed.paragraphs[-1].text == "Firma de Ana"
    assert len(second.inline_shapes) == 0
    assert second.paragraphs[0].text == "Firma de Juan"
