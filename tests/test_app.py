from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document
from streamlit.testing.v1 import AppTest

from test_signing import doc_bytes, drawn_signature


def upload(label="Firma de la persona", name="mi_word.docx"):
    document = Document()
    document.add_paragraph("Documento personal")
    document.add_paragraph(label)
    return SimpleNamespace(name=name, getvalue=lambda: doc_bytes(document))


def test_initial_screen_only_requests_document():
    app = AppTest.from_file("../app.py").run()
    assert not app.exception
    assert app.title[0].value == "Firmar un Word"
    assert not app.subheader
    assert not app.button


def test_signing_and_clearing_invalidates_download():
    pad = SimpleNamespace(signature=None)
    with patch("streamlit.file_uploader", return_value=upload()), patch("signature_pad.signature_pad", return_value=pad):
        app = AppTest.from_file("../app.py").run()
        assert not app.exception
        assert app.button[0].disabled
        pad.signature = drawn_signature()
        app.run()
        assert not app.button[0].disabled
        app.button[0].click().run()
        assert not app.exception
        signed = Document(BytesIO(app.session_state["signed_document"]))
        assert signed.paragraphs[-1].text == "Firma de la persona"
        assert len(signed.inline_shapes) == 1
        pad.signature = None
        app.run()
        assert app.button[0].disabled
        assert "signed_document" not in app.session_state


def test_multiple_locations_require_selection():
    with patch("streamlit.file_uploader", return_value=upload("Firma del docente\n")) as uploader, patch("signature_pad.signature_pad", return_value=SimpleNamespace(signature=None)):
        document = Document()
        document.add_paragraph("Firma de quien recibe")
        document.add_paragraph("Firma de quien entrega")
        uploader.return_value = SimpleNamespace(name="varias.docx", getvalue=lambda: doc_bytes(document))
        app = AppTest.from_file("../app.py").run()
        assert not app.exception
        assert app.selectbox[0].value is None
        assert not app.button
        app.selectbox[0].select(1).run()
        assert not app.exception
        assert app.button[0].disabled


def test_search_without_match_and_invalid_word_are_actionable():
    with patch("streamlit.file_uploader", return_value=upload("Autorización")):
        app = AppTest.from_file("../app.py").run()
        assert not app.exception
        assert "No encontramos" in app.info[0].value
    invalid = SimpleNamespace(name="roto.docx", getvalue=lambda: b"broken")
    with patch("streamlit.file_uploader", return_value=invalid):
        app = AppTest.from_file("../app.py").run()
        assert not app.exception
        assert "No se pudo abrir" in app.error[0].value


def test_sessions_do_not_share_documents_or_signatures():
    with patch("streamlit.file_uploader", return_value=upload("Firma de Ana")), patch("signature_pad.signature_pad", return_value=SimpleNamespace(signature=drawn_signature())):
        first = AppTest.from_file("../app.py").run()
        first.button[0].click().run()
        assert "signed_document" in first.session_state
    with patch("streamlit.file_uploader", return_value=upload("Firma de Juan")), patch("signature_pad.signature_pad", return_value=SimpleNamespace(signature=None)):
        second = AppTest.from_file("../app.py").run()
        assert not second.exception
        assert "signed_document" not in second.session_state
        assert second.button[0].disabled
        assert second.session_state["document"].paragraphs[-1].text == "Firma de Juan"
